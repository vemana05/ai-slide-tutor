import shutil
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from app.config import STORAGE_PATH
from app.parsing.office_render import convert_to_pdf
from app.parsing.pdf_parser import render_pdf_pages


def _render_slide_images(pptx_path: Path, image_dir: Path, slide_count: int) -> list[str | None]:
    """Render each slide to a PNG via headless LibreOffice. Returns None for every
    slide if conversion isn't available, rather than failing the whole upload."""
    try:
        pdf_path = convert_to_pdf(pptx_path)
    except Exception:
        return [None] * slide_count

    try:
        return render_pdf_pages(pdf_path, image_dir, filename_prefix="slide")
    finally:
        shutil.rmtree(pdf_path.parent, ignore_errors=True)


def parse_pptx(pptx_path: Path, image_dir: Path) -> list[dict]:
    """Extract per slide: title/body text, speaker notes, embedded images, and a
    rendered slide image (via headless LibreOffice, so it matches the real layout)."""
    image_dir.mkdir(parents=True, exist_ok=True)
    slides = []

    prs = Presentation(pptx_path)
    rendered_image_paths = _render_slide_images(pptx_path, image_dir, len(prs.slides))

    for slide_index, slide in enumerate(prs.slides):
        text_parts = []
        embedded_image_paths = []
        picture_count = 0

        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                text_parts.append(shape.text_frame.text.strip())

            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                picture_count += 1
                image = shape.image
                image_filename = f"slide_{slide_index + 1}_img_{picture_count}.{image.ext}"
                image_path = image_dir / image_filename
                image_path.write_bytes(image.blob)
                embedded_image_paths.append(image_path.relative_to(STORAGE_PATH).as_posix())

        speaker_notes = None
        if slide.has_notes_slide:
            notes_text = slide.notes_slide.notes_text_frame.text.strip()
            speaker_notes = notes_text or None

        slides.append(
            {
                "raw_text": "\n\n".join(text_parts) or None,
                "speaker_notes": speaker_notes,
                "rendered_image_path": rendered_image_paths[slide_index],
                "embedded_image_paths": embedded_image_paths,
            }
        )

    return slides
